#!/usr/bin/env python3
"""
Extrae los 4 anexos .xlsx → JSONL intermedio.

Uso:
  python extract_annexes.py \\
    --dir /ruta/a/panini_docs/anexos \\
    --out ../data/intermediate

Archivos esperados (nombres exactos):
  lista_primitivos.xlsx
  lista_bases_nominales.xlsx
  pratyahara_sutra.xlsx
  siglas.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Optional

try:
    import openpyxl
except ImportError:
    print("ERROR: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def _cell(row: tuple, idx: int) -> str:
    if idx >= len(row):
        return ""
    v = row[idx]
    if v is None:
        return ""
    return str(v).strip()


def _write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"  → {len(rows)} filas → {path}")


# ---------------------------------------------------------------------------
# 1) lista_primitivos.xlsx
#    Columnas: Grupo | SecciónMayor | Subsección | Nivel4 | Nivel5 | Primitivo/Detalle
# ---------------------------------------------------------------------------

ENTRY_NUM_RE = re.compile(r"^(\d+)\s+(.+)$")


def extract_primitivos(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_out: list[dict[str, Any]] = []
    order = 0

    # Saltar cabecera
    first = True
    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        grupo = _cell(row, 0)
        sec = _cell(row, 1)
        sub = _cell(row, 2)
        n4 = _cell(row, 3)
        n5 = _cell(row, 4)
        detalle = _cell(row, 5)
        if not detalle and not grupo:
            continue

        entry_number: Optional[int] = None
        root = ""
        meaning = ""
        raw = detalle

        # Intentar "935 ayA gat-au" o "1 bhū sattāyām"
        m = ENTRY_NUM_RE.match(detalle)
        if m:
            entry_number = int(m.group(1))
            rest = m.group(2).strip()
            # primera palabra ≈ raíz; resto ≈ significado
            parts = rest.split(None, 1)
            root = parts[0] if parts else rest
            meaning = parts[1] if len(parts) > 1 else ""
        else:
            root = detalle

        rows_out.append(
            {
                "group_label": grupo or None,
                "section_major": sec or None,
                "subsection": sub or None,
                "level4": n4 or None,
                "level5": n5 or None,
                "entry_number": entry_number,
                "root": root,
                "meaning": meaning or None,
                "detail": detalle or None,
                "raw_entry": raw or None,
                "order_index": order,
            }
        )
        order += 1

    wb.close()
    return rows_out


# ---------------------------------------------------------------------------
# 2) lista_bases_nominales.xlsx
#    Columnas: Grupo | Sutra/Referencia | Base Nominal/Detalle
# ---------------------------------------------------------------------------

BASE_NUM_RE = re.compile(r"^(\d+)\s+(.+)$")


def extract_bases(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_out: list[dict[str, Any]] = []
    order = 0
    first = True

    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        grupo = _cell(row, 0)
        sutra_ref = _cell(row, 1)
        base_detail = _cell(row, 2)
        if not base_detail and not grupo:
            continue

        item_number: Optional[int] = None
        base = base_detail
        m = BASE_NUM_RE.match(base_detail)
        if m:
            item_number = int(m.group(1))
            base = m.group(2).strip()

        rows_out.append(
            {
                "group_name": grupo or None,
                "sutra_ref": sutra_ref or None,
                "item_number": item_number,
                "base": base,
                "meaning": None,
                "notes": None,
                "order_index": order,
            }
        )
        order += 1

    wb.close()
    return rows_out


# ---------------------------------------------------------------------------
# 3) pratyahara_sutra.xlsx
#    Columnas: Pratyāhāra-sūtra | Definición | Siglas generadas…
# ---------------------------------------------------------------------------

def extract_pratyahara_sutras(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_out: list[dict[str, Any]] = []
    order = 0
    first = True

    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        # Puede venir con número en col0 o el sūtra en col0
        c0 = _cell(row, 0)
        c1 = _cell(row, 1)
        c2 = _cell(row, 2)
        c3 = _cell(row, 3)

        # Heurística: si c0 es solo dígito, el sūtra está en c1
        if c0.isdigit():
            sutra_text = c1
            definition = c2
            siglas = c3
        else:
            sutra_text = c0
            definition = c1
            siglas = c2

        if not sutra_text:
            continue

        rows_out.append(
            {
                "order_index": order,
                "sutra_text": sutra_text,
                "definition": definition or None,
                "generated_siglas": siglas or None,
            }
        )
        order += 1

    wb.close()
    return rows_out


# ---------------------------------------------------------------------------
# 4) siglas.xlsx  →  se tratan como definiciones de pratyāhāras
#    Columnas: Número | Sigla | Definición/Referencia
# ---------------------------------------------------------------------------

def extract_siglas(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_out: list[dict[str, Any]] = []
    first = True

    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        num = _cell(row, 0)
        sigla = _cell(row, 1)
        definicion = _cell(row, 2)
        if not sigla:
            continue

        order_index = int(num) if num.isdigit() else 0
        # defined_in: extraer primera ref tipo 1.1.3 si aparece
        defined_in = None
        m = re.search(r"\b(\d+\.\d+\.\d+)\b", definicion)
        if m:
            defined_in = m.group(1)

        rows_out.append(
            {
                "id": sigla,
                "expansion": definicion[:200] if definicion else sigla,
                "definition_full": definicion or None,
                "defined_in": defined_in,
                "notes": None,
                "order_index": order_index,
            }
        )

    wb.close()
    return rows_out


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

FILES = {
    "primitivos": "lista_primitivos.xlsx",
    "bases": "lista_bases_nominales.xlsx",
    "pratyahara_sutras": "pratyahara_sutra.xlsx",
    "siglas": "siglas.xlsx",
}


def main() -> int:
    parser = argparse.ArgumentParser(description="Extrae anexos xlsx → JSONL")
    parser.add_argument(
        "--dir",
        required=True,
        help="Directorio con los 4 .xlsx",
    )
    parser.add_argument(
        "--out",
        default="data/intermediate",
        help="Directorio de salida JSONL",
    )
    args = parser.parse_args()

    src = Path(args.dir)
    out = Path(args.out)
    if not src.is_dir():
        print(f"ERROR: no existe el directorio {src}", file=sys.stderr)
        return 1

    mapping = {
        "primitivos": (FILES["primitivos"], extract_primitivos, "dhatus.jsonl"),
        "bases": (FILES["bases"], extract_bases, "nominal_bases.jsonl"),
        "pratyahara_sutras": (
            FILES["pratyahara_sutras"],
            extract_pratyahara_sutras,
            "pratyahara_sutras.jsonl",
        ),
        "siglas": (FILES["siglas"], extract_siglas, "siglas.jsonl"),
    }

    for key, (filename, fn, out_name) in mapping.items():
        path = src / filename
        if not path.exists():
            print(f"⚠ No encontrado: {path} (se omite)")
            continue
        print(f"→ Extrayendo {filename}")
        try:
            rows = fn(path)
        except Exception as exc:
            print(f"ERROR en {filename}: {exc}", file=sys.stderr)
            return 2
        _write_jsonl(out / out_name, rows)

    print("Listo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())